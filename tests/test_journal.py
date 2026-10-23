import json
from pathlib import Path

import pytest

from paint_by_numbers.journal import JournalError, build_journal, write_journal_files


def _order(root: Path, oid, codes, blocked=False, catalog=True):
    d = root / f"order_{oid}"
    (d / "project").mkdir(parents=True)
    (d / "quality").mkdir()
    cat = [{"code": c, "name": c, "hex": "#808080", "lab": [50, 0, 0]} for c in ("A", "B", "C", "D")]
    pal = [{"number": i, "paint_code": c, "paint_name": c, "hex": "#808080", "lab": [50, 0, 0]} for i, c in enumerate(codes, 1)]
    (d / "project" / "project.json").write_text(json.dumps({
        "project_name": f"П{oid}", "created": "2026-10-23T10:00:00+00:00", "source": {"original_name": f"{oid}.jpg"},
        "settings": {"catalog": cat if catalog else []}, "palette": pal, "spec_calculation": {"provisional": True}}), encoding="utf-8")
    (d / "quality" / "quality_report.json").write_text(json.dumps({
        "canvas_mm": [300, 400], "mode": "balanced", "colors_final": len(codes), "regions_final": 50, "blocked": blocked,
        "warnings_summary": {"critical": int(blocked), "warning": 1, "info": 0}, "mean_delta_e76": 7.5,
        "spec_totals": {"paints": len(codes), "containers": 2 * len(codes), "total_ml": 6.0 * len(codes)}}), encoding="utf-8")


def test_journal_and_usage(tmp_path):
    _order(tmp_path, "1", ["A", "B"])
    _order(tmp_path, "2", ["A"])
    _order(tmp_path, "3", ["B", "C"], blocked=True)         # заблоковане не рахується у використанні кольорів
    d = build_journal(tmp_path)
    assert [r["order"] for r in d["orders"]] == ["1", "2", "3"] and d["orders"][2]["status"] == "заблоковано"
    u = {s["paint_code"]: s for s in d["usage"]}
    assert u["A"]["orders"] == 2 and u["B"]["orders"] == 1 and u["A"]["share_pct"] == 100.0
    assert set(d["unused"]) == {"C", "D"} and d["orders_with_catalog"] == 2
    f = write_journal_files(d, tmp_path)
    from openpyxl import load_workbook
    wb = load_workbook(f["xlsx"])
    assert wb.sheetnames == ["Замовлення", "Використання кольорів"] and wb["Замовлення"].cell(4, 12).value == "заблоковано"
    assert f["csv"].read_text(encoding="utf-8-sig").splitlines()[0].startswith("order;name;")


def test_journal_without_catalog_and_errors(tmp_path):
    _order(tmp_path, "1", ["A"], catalog=False)
    d = build_journal(tmp_path)
    assert d["usage"] == [] and d["orders"][0]["catalog"] == "ні"
    with pytest.raises(JournalError):
        build_journal(tmp_path / "порожньо")
