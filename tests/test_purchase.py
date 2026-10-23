import json
import shutil
from pathlib import Path

import pytest

from paint_by_numbers.purchase import PurchaseError, build_purchase_list, write_purchase_files


def _order(root: Path, oid: str, paints, provisional=True):
    d = root / f"order_{oid}"
    (d / "production").mkdir(parents=True)
    (d / "project").mkdir()
    (d / "project" / "project.json").write_text(json.dumps({"project_name": f"П{oid}"}), encoding="utf-8")
    spec = {"order_number": oid, "canvas_mm": [300, 200], "paints": paints,
            "totals": {"containers": sum(p["containers"] for p in paints), "total_ml": sum(p["total_ml"] for p in paints)},
            "calculation": {"provisional": provisional}}
    (d / "production" / "production_specification.json").write_text(json.dumps(spec), encoding="utf-8")


def P(code, cont, lab=(50, 10, 10), de=1.0):
    return {"paint_code": code, "paint_name": f"Фарба {code}", "hex": "#aa5555", "lab": list(lab),
            "area_mm2": 1000.0 * cont, "containers": cont, "total_ml": 3.0 * cont, "catalog_delta_e2000": de}


def test_aggregation(tmp_path):
    _order(tmp_path, "a", [P("S-01", 2), P("S-02", 1)])
    _order(tmp_path, "b", [P("S-01", 3), P("S-10", 4)], provisional=False)
    d = build_purchase_list(tmp_path, bulk_ml=10, fill_loss_pct=10)
    by = {e["paint_code"]: e for e in d["paints"]}
    assert by["S-01"]["containers"] == 5 and by["S-01"]["orders"] == {"a": 2, "b": 3}
    assert [e["paint_code"] for e in d["paints"]] == ["S-01", "S-02", "S-10"]   # природний порядок
    assert d["totals"]["containers"] == 10 and d["provisional"] and d["provisional_orders"] == ["a"]
    assert by["S-01"]["bulk_units"] == 2        # 15 мл * 1.1 / 10 = 1.65 -> 2


def test_excludes_orders_without_catalog(tmp_path):
    _order(tmp_path, "a", [P("S-01", 2)])
    _order(tmp_path, "b", [P("P-001", 2, de=None)])
    d = build_purchase_list(tmp_path)
    assert [o["order"] for o in d["orders"]] == ["a"]
    assert d["excluded_orders"][0]["order"] == "b"


def test_no_catalog_at_all(tmp_path):
    _order(tmp_path, "a", [P("P-001", 2, de=None)])
    with pytest.raises(PurchaseError):
        build_purchase_list(tmp_path)


def test_conflict_and_errors(tmp_path):
    _order(tmp_path, "a", [P("S-01", 1)])
    _order(tmp_path, "b", [P("S-01", 1, lab=(80, -20, 30))])
    assert build_purchase_list(tmp_path)["code_conflicts"] == [{"paint_code": "S-01", "order": "b"}]
    with pytest.raises(PurchaseError):
        build_purchase_list(tmp_path, bulk_ml=0)
    with pytest.raises(PurchaseError):
        build_purchase_list(tmp_path / "нема")


def test_files(tmp_path):
    _order(tmp_path, "a", [P("S-01", 2)])
    _order(tmp_path, "b", [P("S-01", 1), P("S-02", 1)])
    f = write_purchase_files(build_purchase_list(tmp_path, bulk_ml=20), tmp_path)
    assert all(p.stat().st_size > 0 for p in f.values())
    from openpyxl import load_workbook
    wb = load_workbook(f["xlsx"])
    assert wb.sheetnames[:2] == ["Закупівля", "По замовленнях"]
    ws = wb["Закупівля"]
    assert ws.cell(1, 1).value.startswith("ОРІЄНТОВНИЙ")
    assert ws.cell(4, 1).value == "S-01" and ws.cell(4, 6).value == 3
    assert str(ws.cell(6, 6).value).startswith("=SUM")
    rows = f["csv"].read_text(encoding="utf-8-sig").splitlines()
    assert rows[0].startswith("paint_code;") and rows[-1].startswith("TOTAL;")
    assert f["pdf"].read_bytes()[:4] == b"%PDF"
