"""Журнал замовлень: один реєстр усіх пакетів у папці та статистика використання кольорів складу.

Джерело даних --- самі пакети (project/project.json, quality/quality_report.json, production/…json), окремої бази
немає, тож журнал завжди можна перебудувати з нуля. Статистика допомагає вирішити, які кольори складу
використовуються рідко (кандидати на вилучення) і де похибка велика (кандидати на додавання)."""
from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

from .batch import natural_key

COLUMNS = ["order", "name", "created", "photo", "canvas_mm", "mode", "colors", "regions", "containers", "total_ml",
           "catalog", "status", "critical", "warnings", "provisional", "mean_delta_e76"]


class JournalError(ValueError):
    pass


def _read(path: Path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def build_journal(out_root: str | Path, stock: list[dict] | None = None) -> dict:
    root = Path(out_root)
    if not root.is_dir():
        raise JournalError(f"«{root}» не є папкою.")
    dirs = sorted((d for d in root.iterdir() if d.is_dir() and d.name.startswith("order_")
                   and (d / "project" / "project.json").is_file()), key=lambda d: natural_key(d.name))
    if not dirs:
        raise JournalError("Немає замовлень (папок order_* із project/project.json).")
    rows, usage, stock_from = [], {}, None
    for d in dirs:
        pj = _read(d / "project" / "project.json") or {}
        qr = _read(d / "quality" / "quality_report.json") or {}
        spec = _read(d / "production" / "production_specification.json")
        sm = qr.get("warnings_summary", {})
        tot = (spec or {}).get("totals") or qr.get("spec_totals") or {}
        cat = pj.get("settings", {}).get("catalog") or []
        used_catalog = bool(cat)
        if used_catalog:
            stock_from = cat
        status = "заблоковано" if qr.get("blocked") else "готово"
        rows.append({
            "order": d.name.removeprefix("order_"), "name": pj.get("project_name", ""),
            "created": (pj.get("created") or "")[:19].replace("T", " "), "photo": pj.get("source", {}).get("original_name", ""),
            "canvas_mm": "x".join(f"{v:g}" for v in qr.get("canvas_mm", [])), "mode": qr.get("mode", ""),
            "colors": tot.get("paints", qr.get("colors_final", "")), "regions": qr.get("regions_final", ""),
            "containers": tot.get("containers", ""), "total_ml": tot.get("total_ml", ""),
            "catalog": "так" if used_catalog else "ні", "status": status, "critical": sm.get("critical", ""),
            "warnings": sm.get("warning", ""),
            "provisional": "так" if (pj.get("spec_calculation") or {}).get("provisional", True) else "ні",
            "mean_delta_e76": qr.get("mean_delta_e76", "")})
        if used_catalog and not qr.get("blocked"):
            cont = {p["paint_code"]: p["containers"] for p in (spec or {}).get("paints", [])}
            for p in pj.get("palette", []):
                u = usage.setdefault(p["paint_code"], {"paint_code": p["paint_code"], "paint_name": p.get("paint_name", ""),
                                                       "hex": p.get("hex", ""), "orders": 0, "containers": 0})
                u["orders"] += 1
                u["containers"] += cont.get(p["paint_code"], 0)
    stock = stock or stock_from or []
    n_cat = sum(1 for r in rows if r["catalog"] == "так" and r["status"] == "готово")
    stats = []
    for c in stock:
        u = usage.get(c["code"], {})
        stats.append({"paint_code": c["code"], "paint_name": c.get("name", ""), "hex": c.get("hex", ""),
                      "orders": u.get("orders", 0), "share_pct": round(100.0 * u.get("orders", 0) / n_cat, 1) if n_cat else 0.0,
                      "containers": u.get("containers", 0)})
    known = {c["code"] for c in stock}
    for code, u in usage.items():             # коди, яких немає в поточному складі (змінений каталог)
        if code not in known:
            stats.append({"paint_code": code, "paint_name": u["paint_name"], "hex": u["hex"], "orders": u["orders"],
                          "share_pct": round(100.0 * u["orders"] / n_cat, 1) if n_cat else 0.0,
                          "containers": u["containers"], "not_in_stock": True})
    stats.sort(key=lambda s: (-s["orders"], natural_key(s["paint_code"])))
    return {"created": datetime.now().isoformat(timespec="seconds"), "orders": rows, "usage": stats,
            "orders_with_catalog": n_cat, "unused": [s["paint_code"] for s in stats if s["orders"] == 0]}


def write_journal_files(data: dict, out_dir: str | Path) -> dict[str, Path]:
    d = Path(out_dir)
    files = {"csv": d / "orders_journal.csv", "xlsx": d / "orders_journal.xlsx", "json": d / "orders_journal.json"}
    with open(files["csv"], "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(COLUMNS)
        for r in data["orders"]:
            w.writerow([r[c] for c in COLUMNS])
    files["json"].write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Замовлення"
    head, bold = PatternFill("solid", fgColor="E6ECF5"), Font(bold=True)
    titles = ["Замовлення", "Назва", "Створено", "Фото", "Полотно, мм", "Режим", "Кольорів", "Ділянок", "Ємностей", "мл",
              "Каталог", "Статус", "Критичних", "Попереджень", "Орієнтовно", "ΔE76"]
    for j, t in enumerate(titles, 1):
        c = ws.cell(1, j, t)
        c.font, c.fill, c.alignment = bold, head, Alignment(horizontal="center", wrap_text=True)
    for i, r in enumerate(data["orders"], 2):
        for j, k in enumerate(COLUMNS, 1):
            ws.cell(i, j, r[k])
        if r["status"] != "готово":
            ws.cell(i, 12).font = Font(bold=True, color="C00000")
    for j, w in enumerate([13, 18, 19, 22, 12, 11, 9, 9, 10, 8, 9, 13, 10, 12, 11, 8], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(data['orders']) + 1}"
    u = wb.create_sheet("Використання кольорів")
    for j, t in enumerate(["Код", "Назва", "Колір", "HEX", "Замовлень", "% замовлень", "Ємностей", "Примітка"], 1):
        c = u.cell(1, j, t)
        c.font, c.fill = bold, head
    for i, s in enumerate(data["usage"], 2):
        u.cell(i, 1, s["paint_code"])
        u.cell(i, 2, s["paint_name"])
        if s["hex"]:
            u.cell(i, 3).fill = PatternFill("solid", fgColor=s["hex"].lstrip("#"))
        u.cell(i, 4, s["hex"])
        u.cell(i, 5, s["orders"])
        u.cell(i, 6, s["share_pct"])
        u.cell(i, 7, s["containers"])
        u.cell(i, 8, "немає в поточному складі" if s.get("not_in_stock") else ("не використовувався" if s["orders"] == 0 else ""))
    for j, w in enumerate([10, 26, 8, 10, 11, 12, 10, 28], 1):
        u.column_dimensions[get_column_letter(j)].width = w
    u.freeze_panes = "A2"
    wb.save(files["xlsx"])
    return files
