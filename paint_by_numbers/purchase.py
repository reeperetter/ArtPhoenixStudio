"""Загальний список закупівлі фарб для партії замовлень (розділ 11 ТЗ): скільки ємностей кожної фарби
потрібно підготувати на ВСІ замовлення разом.

Фарби різних замовлень можна складати лише тоді, коли це фарби з каталогу (однаковий `paint_code` = та сама
фарба). Замовлення без каталогу мають унікальні кольори, які не збігаються між замовленнями, тому в загальний
список не входять, а перелічуються окремо з причиною. Ємності кожного замовлення округлюються вгору
окремо (кожен набір отримує власні ємності), тож загальна кількість = сума по замовленнях."""
from __future__ import annotations

import csv
import json
import math
from datetime import datetime
from pathlib import Path

import numpy as np

from .batch import natural_key
from .palette.catalog import ciede2000

LAB_CONFLICT_DE = 1.0       # однаковий код із відмінним кольором (ΔE2000 більше) --- ознака змішаних каталогів


class PurchaseError(ValueError):
    pass


def _read(path: Path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def collect_orders(out_root: str | Path) -> list[Path]:
    root = Path(out_root)
    if not root.is_dir():
        raise PurchaseError(f"«{root}» не є папкою.")
    return sorted((d for d in root.iterdir() if d.is_dir() and d.name.startswith("order_")
                   and (d / "production" / "production_specification.json").is_file()),
                  key=lambda d: natural_key(d.name))


def build_purchase_list(out_root: str | Path, bulk_ml: float | None = None, fill_loss_pct: float = 0.0,
                        orders: list[str] | None = None) -> dict:
    """Читає готові замовлення з `out_root` (папки order_*). `orders` --- необовʼязковий відбір за номерами."""
    if bulk_ml is not None and bulk_ml <= 0:
        raise PurchaseError("Обʼєм великої тари має бути додатним.")
    if not 0 <= fill_loss_pct < 100:
        raise PurchaseError("Втрати при розливі мають бути в межах 0..100 %.")
    dirs = collect_orders(out_root)
    if orders is not None:
        want = {o.removeprefix("order_") for o in orders}
        dirs = [d for d in dirs if d.name.removeprefix("order_") in want]
    if not dirs:
        raise PurchaseError("Немає готових замовлень (папок order_* із production/production_specification.json).")

    paints: dict[str, dict] = {}
    included, excluded, conflicts = [], [], []
    provisional_orders = []
    for d in dirs:
        oid = d.name.removeprefix("order_")
        spec = _read(d / "production" / "production_specification.json")
        meta = _read(d / "project" / "project.json") or {}
        if not spec or "paints" not in spec:
            excluded.append({"order": oid, "reason": "пошкоджена специфікація"})
            continue
        if not all(p.get("catalog_delta_e2000") is not None for p in spec["paints"]):
            excluded.append({"order": oid, "reason": "без каталогу фарб: кольори замовлення унікальні "
                                                     "(повторіть генерацію з --catalog)"})
            continue
        name = meta.get("project_name", "")
        calc = spec.get("calculation", {})
        prov = bool(calc.get("provisional", True))
        if prov:
            provisional_orders.append(oid)
        included.append({"order": oid, "name": name, "containers": spec["totals"]["containers"],
                         "total_ml": spec["totals"]["total_ml"], "provisional": prov,
                         "canvas_mm": spec.get("canvas_mm")})
        for p in spec["paints"]:
            e = paints.setdefault(p["paint_code"], {
                "paint_code": p["paint_code"], "paint_name": p["paint_name"], "hex": p["hex"], "lab": p["lab"],
                "orders": {}, "containers": 0, "total_ml": 0.0, "area_mm2": 0.0})
            if ciede2000(np.array(e["lab"]), np.array(p["lab"])) > LAB_CONFLICT_DE:
                conflicts.append({"paint_code": p["paint_code"], "order": oid})
            e["orders"][oid] = e["orders"].get(oid, 0) + p["containers"]
            e["containers"] += p["containers"]
            e["total_ml"] = round(e["total_ml"] + p["total_ml"], 2)
            e["area_mm2"] = round(e["area_mm2"] + p["area_mm2"], 1)
    if not included:
        raise PurchaseError("У жодному замовленні не використано каталог фарб: загальний список неможливий. "
                            "Повторіть генерацію з каталогом (--catalog).")
    rows = sorted(paints.values(), key=lambda e: natural_key(e["paint_code"]))
    for e in rows:
        e["orders_count"] = len(e["orders"])
        if bulk_ml:
            e["bulk_units"] = math.ceil(round(e["total_ml"] * (1 + fill_loss_pct / 100.0) / bulk_ml, 9))
    totals = {"orders": len(included), "paints": len(rows), "containers": sum(e["containers"] for e in rows),
              "total_ml": round(sum(e["total_ml"] for e in rows), 2),
              "area_mm2": round(sum(e["area_mm2"] for e in rows), 1)}
    if bulk_ml:
        totals["bulk_units"] = sum(e["bulk_units"] for e in rows)
    return {
        "created": datetime.now().isoformat(timespec="seconds"), "root": str(out_root), "paints": rows,
        "totals": totals, "orders": included, "excluded_orders": excluded, "code_conflicts": conflicts,
        "provisional": bool(provisional_orders), "provisional_orders": provisional_orders,
        "bulk_ml": bulk_ml, "fill_loss_pct": fill_loss_pct if bulk_ml else None,
    }


# ------------------------------------------------------------------ файли
def write_csv(data: dict, path: str | Path) -> None:
    cols = ["paint_code", "paint_name", "hex", "orders_count", "containers", "total_ml", "area_mm2"]
    if data["bulk_ml"]:
        cols.append("bulk_units")
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(cols + ["estimate"])
        est = "orientational" if data["provisional"] else "confirmed"
        for e in data["paints"]:
            w.writerow([e.get(c, "") for c in cols] + [est])
        t = data["totals"]
        w.writerow(["TOTAL", "", "", t["orders"], t["containers"], t["total_ml"], t["area_mm2"]]
                   + ([t["bulk_units"]] if data["bulk_ml"] else []) + [""])


def write_xlsx(data: dict, path: str | Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Закупівля"
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="E6ECF5")
    cols = ["Код фарби", "Назва", "Колір", "HEX", "Замовлень", "Ємностей", "Обʼєм, мл", "Площа, мм²"]
    if data["bulk_ml"]:
        cols.append(f"Тара {data['bulk_ml']:g} мл, шт")
    r0 = 1
    if data["provisional"]:
        ws.cell(1, 1, "ОРІЄНТОВНИЙ РОЗРАХУНОК: норма покриття не підтверджена фізичними тестами").font = Font(bold=True, color="C00000")
        r0 = 3
    for j, c in enumerate(cols, 1):
        cell = ws.cell(r0, j, c)
        cell.font, cell.fill = bold, head
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, e in enumerate(data["paints"], r0 + 1):
        ws.cell(i, 1, e["paint_code"])
        ws.cell(i, 2, e["paint_name"])
        ws.cell(i, 3).fill = PatternFill("solid", fgColor=e["hex"].lstrip("#"))
        ws.cell(i, 4, e["hex"])
        ws.cell(i, 5, e["orders_count"])
        ws.cell(i, 6, e["containers"])
        ws.cell(i, 7, e["total_ml"])
        ws.cell(i, 8, e["area_mm2"])
        if data["bulk_ml"]:
            ws.cell(i, 9, e["bulk_units"])
    first, last = r0 + 1, r0 + len(data["paints"])
    tr = last + 1
    ws.cell(tr, 1, "РАЗОМ").font = bold
    for j in range(6, len(cols) + 1):
        L = get_column_letter(j)
        ws.cell(tr, j, f"=SUM({L}{first}:{L}{last})").font = bold
    for j, w in enumerate([14, 28, 8, 10, 11, 11, 11, 13, 14], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(first, 1)

    # матриця «фарба x замовлення»
    m = wb.create_sheet("По замовленнях")
    oids = [o["order"] for o in data["orders"]]
    m.cell(1, 1, "Код фарби").font = bold
    m.cell(1, 2, "Назва").font = bold
    for j, o in enumerate(data["orders"], 3):
        c = m.cell(1, j, o["order"] + (f"\n{o['name']}" if o["name"] and o["name"] != o["order"] else ""))
        c.font, c.alignment = bold, Alignment(wrap_text=True, horizontal="center")
    m.cell(1, len(oids) + 3, "Разом").font = bold
    for i, e in enumerate(data["paints"], 2):
        m.cell(i, 1, e["paint_code"])
        m.cell(i, 2, e["paint_name"])
        for j, o in enumerate(oids, 3):
            if o in e["orders"]:
                m.cell(i, j, e["orders"][o])
        L1, L2 = get_column_letter(3), get_column_letter(len(oids) + 2)
        m.cell(i, len(oids) + 3, f"=SUM({L1}{i}:{L2}{i})").font = bold
    n = len(data["paints"]) + 2
    m.cell(n, 1, "РАЗОМ").font = bold
    for j in range(3, len(oids) + 4):
        L = get_column_letter(j)
        m.cell(n, j, f"=SUM({L}2:{L}{n - 1})").font = bold
    m.column_dimensions["A"].width, m.column_dimensions["B"].width = 14, 28
    m.freeze_panes = "C2"
    if data["excluded_orders"] or data["code_conflicts"]:
        x = wb.create_sheet("Увага")
        x.cell(1, 1, "Замовлення").font = bold
        x.cell(1, 2, "Причина").font = bold
        r = 2
        for e in data["excluded_orders"]:
            x.cell(r, 1, e["order"])
            x.cell(r, 2, "не включено: " + e["reason"])
            r += 1
        for c in data["code_conflicts"]:
            x.cell(r, 1, c["order"])
            x.cell(r, 2, f"код {c['paint_code']} має інший колір, ніж в інших замовленнях (змішані каталоги?)")
            r += 1
        x.column_dimensions["A"].width, x.column_dimensions["B"].width = 16, 90
    wb.save(path)


def write_pdf(data: dict, path: str | Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from .export.pdf import _register_font
    font, bold, cyr = _register_font()
    uk = cyr
    T = (dict(title="Список закупівлі фарб", orders="Замовлень", paints="Фарб", cont="Ємностей", ml="мл",
              code="Код", name="Назва", color="Колір", n="Замовл.", c="Ємн.", v="Обʼєм, мл", total="Разом",
              prov="ОРІЄНТОВНИЙ РОЗРАХУНОК: норма покриття не підтверджена фізичними тестами.",
              bulk="Тара {b:g} мл, шт", excl="Не включено", conf="Увага: код має різні кольори в замовленнях",
              date="Дата") if uk else
         dict(title="Paint purchase list", orders="Orders", paints="Paints", cont="Containers", ml="ml",
              code="Code", name="Name", color="Color", n="Orders", c="Cont.", v="Volume, ml", total="Total",
              prov="PROVISIONAL CALCULATION: coverage norm not confirmed by physical tests.",
              bulk="Bulk {b:g} ml, pcs", excl="Excluded", conf="Warning: code has different colors across orders",
              date="Date"))
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm,
                            bottomMargin=15 * mm, title=T["title"], author="paint-by-numbers")
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=16, spaceAfter=4)
    nm = ParagraphStyle("nm", fontName=font, fontSize=9.5, leading=13)
    warn = ParagraphStyle("w", fontName=bold, fontSize=10, textColor=colors.HexColor("#C00000"), spaceAfter=4)
    t = data["totals"]
    flow = [Paragraph(T["title"], h1),
            Paragraph(f'{T["date"]}: {data["created"][:10]}. {T["orders"]}: {t["orders"]}; {T["paints"]}: {t["paints"]}; '
                      f'{T["cont"]}: {t["containers"]} ({t["total_ml"]:g} {T["ml"]})', nm)]
    if data["provisional"]:
        flow.append(Paragraph(T["prov"], warn))
    flow.append(Spacer(1, 4 * mm))
    head = [T["code"], T["name"], T["color"], T["n"], T["c"], T["v"]]
    if data["bulk_ml"]:
        head.append(T["bulk"].format(b=data["bulk_ml"]))
    rows = [head]
    for e in data["paints"]:
        rows.append([e["paint_code"], Paragraph(e["paint_name"] or "-", nm), "", e["orders_count"], e["containers"],
                     f'{e["total_ml"]:g}'] + ([e["bulk_units"]] if data["bulk_ml"] else []))
    rows.append([T["total"], "", "", t["orders"], t["containers"], f'{t["total_ml"]:g}']
                + ([t["bulk_units"]] if data["bulk_ml"] else []))
    widths = [24, 60, 16, 18, 18, 24] + ([26] if data["bulk_ml"] else [])
    tbl = Table(rows, colWidths=[w * mm for w in widths], repeatRows=1)
    st = [("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 9),
          ("FONTNAME", (0, 0), (-1, 0), bold), ("FONTNAME", (0, -1), (-1, -1), bold),
          ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6ECF5")),
          ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#888888")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("ALIGN", (3, 0), (-1, -1), "CENTER")]
    for i, e in enumerate(data["paints"], 1):
        st.append(("BACKGROUND", (2, i), (2, i), colors.HexColor(e["hex"])))
    tbl.setStyle(TableStyle(st))
    flow.append(tbl)
    if data["excluded_orders"] or data["code_conflicts"]:
        flow.append(Spacer(1, 5 * mm))
        for e in data["excluded_orders"]:
            flow.append(Paragraph(f'{T["excl"]}: {e["order"]} --- {e["reason"]}', nm))
        for c in data["code_conflicts"]:
            flow.append(Paragraph(f'{T["conf"]}: {c["paint_code"]} ({c["order"]})', nm))
    doc.build(flow)


def write_purchase_files(data: dict, out_dir: str | Path) -> dict[str, Path]:
    d = Path(out_dir)
    d.mkdir(parents=True, exist_ok=True)
    files = {"csv": d / "purchase_list.csv", "xlsx": d / "purchase_list.xlsx", "pdf": d / "purchase_list.pdf",
             "json": d / "purchase_list.json"}
    write_csv(data, files["csv"])
    write_xlsx(data, files["xlsx"])
    write_pdf(data, files["pdf"])
    files["json"].write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return files
